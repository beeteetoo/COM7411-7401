"""Build concise shared Research Methods decks from original teaching copy.

Requires python-pptx. Source chapters are staff references; no book text or figures
are copied into the student-facing decks. Run from any working directory.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Slides"
OUT.mkdir(exist_ok=True)

NAVY = RGBColor(18, 32, 54)
BLUE = RGBColor(24, 83, 117)
TEAL = RGBColor(0, 154, 149)
WHITE = RGBColor(255, 255, 255)
PALE = RGBColor(226, 240, 244)
AMBER = RGBColor(246, 184, 73)


SPECS = [
    dict(n="01", title="Foundations and research quality", lo="LO1", portfolio="Background & Context", source="Chapters 1–3, pp. 1–36",
         ideas=["Research asks a question whose answer is not yet established.", "Connect the problem, existing knowledge, gap and evidence.", "A prototype can be a method; building it alone is not a finding."],
         example=["A browser warning is proposed for suspicious login pages.", "Compare recognition accuracy with an existing warning in a defined group.", "The sample and evaluation design bound any claim of improvement."],
         checks=["What is still uncertain?", "Which observation could answer the question?", "What claim would the proposed evidence not support?"],
         task="Write a 120-word provisional problem, significance and evidence statement.", success="Name a bounded uncertainty and plausible evidence.", terms="Problem • evidence • feasibility • artefact"),
    dict(n="02", title="Research philosophy", lo="LO1", portfolio="Critical Literature Review", source="Chapter 6, pp. 81–96",
         ideas=["Philosophical ontology concerns the nature of what is studied.", "Epistemology concerns what could count as knowledge.", "A computing domain ontology is a different technical concept."],
         example=["Dashboard use could be studied through task accuracy and time.", "It could also be studied through analysts’ interpretations in context.", "Choose a position for the question and evaluate a credible alternative."],
         checks=["What evidence does your question need?", "Which assumption makes that evidence meaningful?", "What would a serious alternative reveal?"],
         task="State your question, chosen position, evidence and counter-position in four sentences.", success="Justify the study-specific choice without assigning a paradigm to a whole discipline.", terms="Ontology • epistemology • paradigm • counter-position"),
    dict(n="03", title="From topic to research question", lo="LO1–2", portfolio="Background & Context", source="Chapter 4, pp. 39–57; Chapter 3, pp. 29–31",
         ideas=["Funnel: field → evidence → gap → scoped problem → question.", "A gap needs a defensible source trail, not an assertion of novelty.", "Design-and-creation research still needs an evaluative knowledge claim."],
         example=["‘Build an AI app’ states a task.", "‘How accurately does it identify support requests versus a baseline?’ asks for evidence.", "Define the dataset, outcome measure, comparison and limitation."],
         checks=["Which key word needs defining?", "Can the proposed evidence answer the question?", "Does it ask for knowledge or only build steps?"],
         task="Create a six-box funnel, write one question and a two-sentence evidence plan.", success="The question is specific, answerable and feasible.", terms="Scope • gap • baseline • operational definition"),
    dict(n="04", title="Critical literature review", lo="LO1–3", portfolio="Critical Literature Review", source="Chapter 5, pp. 59–79",
         ideas=["Record each source's question, method, evidence, claim and limit.", "Compare sources to explain agreement, conflict and uncertainty.", "End with the gap and why your proposal follows from it."],
         example=["An experiment reports better phishing-test scores.", "A case study finds warnings are ignored under time pressure.", "A design paper proposes an interface but has not established effectiveness."],
         checks=["What can each study actually establish?", "How do at least two sources relate?", "What uncertainty remains for your question?"],
         task="Build a three-source matrix and a 180-word comparative paragraph.", success="Compare evidence and end with a justified gap.", terms="Synthesis • appraisal • grey literature • search trail"),
    dict(n="05", title="Research design and ethics", lo="LO1–2", portfolio="Proposed Methodology", source="Chapter 6, pp. 97–112; Chapter 7, pp. 125–127",
         ideas=["Connect question → evidence → design → method → analysis.", "An artefact needs evaluation; mixed methods need an integration step.", "Consent, harm, privacy and access can change the design itself."],
         example=["Click logs could describe warning use; interviews could explain context.", "The proposal states where the two strands will be combined.", "It also specifies data access, consent and protection."],
         checks=["What would the method let you conclude?", "Which alternative did you reject and why?", "Which safeguard changes the procedure?"],
         task="Draw a design chain and name one specific risk, mitigation and rejected alternative.", success="Every link serves the question and is feasible and ethical.", terms="Design • method • integration • mitigation"),
    dict(n="06", title="Sampling and data collection", lo="LO2", portfolio="Proposed Methodology", source="Chapter 7, pp. 115–129",
         ideas=["Define the unit of analysis and the source or population.", "State a selection rule and an instrument or procedure.", "Match the evidence to the question and protect access and people."],
         example=["To compare alert interfaces, use simulated alerts and a defined task.", "Specify analyst eligibility, outcome measures and order effects.", "A convenient small sample does not represent all analysts."],
         checks=["What does one row or account represent?", "How would cases be selected?", "What data and permission are needed?"],
         task="Complete the five-row collection plan and draft one item or measure.", success="Source, selection and instrument form a feasible chain.", terms="Unit of analysis • sampling frame • instrument • secondary data"),
    dict(n="07", title="Quantitative analysis plan", lo="LO2–3", portfolio="Proposed Methodology", source="Chapter 7, pp. 130–137; Chapter 6, pp. 107–109",
         ideas=["Define each variable and the observation it describes.", "Choose a statistic to answer the question, not because software offers it.", "Separate description, association and causal inference."],
         example=["The invented CSV compares five attempts per warning interface.", "Calculate correct-flag counts and median response time.", "These ten cases cannot establish a population effect or causation."],
         checks=["Which measure fits the question?", "What assumption does the comparison need?", "What cannot be concluded?"],
         task="Analyse the invented CSV and name the analysis planned for your proposal.", success="Method and interpretation fit the data and limits.", terms="Variable • median • association • inference"),
    dict(n="08", title="Qualitative analysis plan", lo="LO2–3", portfolio="Proposed Methodology", source="Chapter 6, pp. 109–111; Chapter 7, pp. 133–135",
         ideas=["Move from meaningful text segments to codes and candidate themes.", "A theme is an interpretive claim, not a frequent word.", "Keep a traceable path and seek a rival interpretation."],
         example=["‘Too many warnings’ can be coded alert overload.", "‘I ask a colleague’ can be coded social verification.", "A possible theme links warning use to workflow and trust."],
         checks=["Which quotation supports the code?", "How does the theme answer the question?", "What other reading fits the extract?"],
         task="Code four lines of the invented transcript; propose a theme and rival reading.", success="A supported interpretation and method choice, with a rigour step.", terms="Code • theme • audit trail • reflexivity"),
    dict(n="09", title="Proposal integration and reflection", lo="LO1–3", portfolio="Full report", source="Chapters 3 and 5; Chapter 8, pp. 139–149",
         ideas=["The four report sections should form one argument.", "Literature justifies the method; the method could answer the question.", "Reflection names a real trade-off and suitable quality check."],
         example=["A gap in click-rate studies may justify contextual interviews.", "A small purposive sample offers depth, not a population estimate.", "An audit trail and disconfirming cases could strengthen rigour."],
         checks=["Is the same question used throughout?", "Could the data answer it?", "Is the limitation specific to this design?"],
         task="Peer-review a draft, record two revisions and improve the reflection before report lock.", success="A coherent, ethical and feasible report ready for submission.", terms="Coherence • limitation • rigour • reflection"),
    dict(n="10", title="Four-minute methodological defence", lo="LO1–3", portfolio="Recorded oral defence", source="Chapter 3, pp. 34–36; Chapter 8, pp. 149–153",
         ideas=["State the question exactly as submitted in the locked report.", "Name and apply the assigned methodological check directly.", "Explain a bounded conclusion within four minutes ±30 seconds."],
         example=["A prototype might improve a specified decision.", "A defence explains the comparison and measure that could test this.", "The build itself is not proof of improvement."],
         checks=["Did you name the assigned check?", "Did you apply it to your own question and method?", "Did you explain a limit rather than recite the report?"],
         task="Practise with a sample check; record the assigned-check response after report lock.", success="Clear unaided reasoning in the approved format.", terms="Defence • alignment • counterargument • limitation"),
]


def box(slide, x, y, w, h, fill, text, size=24, color=WHITE, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.22)
    frame.margin_right = Inches(0.18)
    frame.margin_top = Inches(0.12)
    frame.margin_bottom = Inches(0.08)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def base_slide(prs, label, source):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    box(slide, 0.55, 0.35, 12.2, 0.72, BLUE, label, 28, WHITE, True)
    footer = slide.shapes.add_textbox(Inches(0.65), Inches(7.07), Inches(12.0), Inches(0.26))
    p = footer.text_frame.paragraphs[0]
    p.text = f"Research Methods • shared academic master • Mbanaso et al. (2023), {source}"
    p.font.name = "Aptos"
    p.font.size = Pt(11)
    p.font.color.rgb = PALE
    return slide


def three_rows(slide, items, accent=TEAL):
    for i, item in enumerate(items):
        box(slide, 0.9, 1.45 + i * 1.65, 11.5, 1.30, accent if i == 0 else BLUE, item, 25)


for spec in SPECS:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    title = base_slide(prs, spec["title"], spec["source"])
    box(title, 1.0, 1.55, 11.3, 1.4, TEAL, f"Topic {spec['n']} • {spec['lo']}", 32, WHITE, True)
    box(title, 1.0, 3.35, 11.3, 1.45, BLUE, f"Portfolio link: {spec['portfolio']}", 27)
    box(title, 1.0, 5.25, 11.3, 0.85, NAVY, "One shared concept set • track-specific live activity", 23, PALE)

    s = base_slide(prs, "Three ideas to take away", spec["source"])
    three_rows(s, spec["ideas"])

    s = base_slide(prs, "Worked computing example", spec["source"])
    three_rows(s, spec["example"], BLUE)

    s = base_slide(prs, "Pause and test the reasoning", spec["source"])
    three_rows(s, spec["checks"])

    s = base_slide(prs, "Your individual task", spec["source"])
    box(s, 0.9, 1.55, 11.5, 2.2, BLUE, spec["task"], 29)
    box(s, 0.9, 4.1, 11.5, 1.5, TEAL, "Success check: " + spec["success"], 25)

    s = base_slide(prs, "Take it into the portfolio", spec["source"])
    box(s, 0.9, 1.55, 11.5, 1.55, BLUE, f"Report connection: {spec['portfolio']}", 27)
    box(s, 0.9, 3.4, 11.5, 1.55, TEAL, "Key terms: " + spec["terms"], 24)
    box(s, 0.9, 5.2, 11.5, 0.75, NAVY, "Use the numbered student note for the full explanation and task.", 21, PALE)

    out = OUT / f"{spec['n']} {spec['title']}.pptx"
    prs.save(out)
    print(out)
